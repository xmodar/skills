"""The instructor's per-round table, and the "what if the Present threshold were lower" simulation.

    python -I score_table.py table   <course_dir> <workbook.xlsx> <grades_dir> <out.md>
    python -I score_table.py whatif  <course_dir> <workbook.xlsx> <grades_dir> <out.txt> --present-min 0.40 [--late-min 0.25]

table: one row per submission in submissions.json, ORDERED BY POSTING TIME:
    | # | chat name | project score | remaining | notes |
  - project score = min(100, A..E + max(grader bonus, bonus review)), bonus capped at grading.bonus_max (10); a student
    with several submissions gets the best one (only the best is listed);
  - remaining = 0 if the certificate is اجتياز (total > 65 and attendance > 65), otherwise the project points still
    needed: floor((65 − 0.3 × attendance) / 0.7) + 1 − project; "—" when attendance ≤ 65% (no project score can pass);
  - notes = empty when passing; the grader's "missing" keywords when short on project; "attendance N%" when blocked.
  Prints: submitted N | pass N.
whatif: recomputes ClassPulse days from each learner's answered/total checkpoints in classpulse_marks.json with the given
  thresholds (chat days and status_overrides kept) and lists everyone whose marks change, with attendance and
  certificate before → after; certificate changes first. Nothing is written to the workbook.
"""
import argparse, json, math, os, re
import openpyxl


def load(p, d=None):
    return json.load(open(p, encoding="utf-8")) if os.path.exists(p) else d


def projects(D, G):
    subs, st = load(os.path.join(D, "submissions.json"), {}), load(os.path.join(D, "submissions_state.json"), {})
    bmax = (load(os.path.join(D, "course.json"), {}).get("grading") or {}).get("bonus_max", 10)
    best = {}
    for d, x in subs.items():
        v = (st.get(d) or {}).get("version", 1)
        g = load(os.path.join(G, f"{d}.v{v}.json"))
        if not g:
            continue
        b = (load(os.path.join(G, f"{d}.v{v}.bonus.json")) or {}).get("bonus", 0)
        letters = [k for k in "ABCDEFGHIJ" if k in g]
        P = min(100, sum(int(g[k]) for k in letters) + min(bmax, max(int(g.get("bonus") or 0), int(b))))
        e = str(x["student"]).lower()
        if e not in best or P > best[e][0]:
            best[e] = (P, d, x, g)
    return best


SECTION = re.compile(r'\[(>=|<=|>|<|=)\s*([\d.]+)\]\s*"([^"]+)"')


def mark(cell):
    """Text mark, or a ClassPulse ratio formula ("=4/6") read through its conditional number format."""
    v = cell.value
    m = re.fullmatch(r"=\s*(\d+(?:\.\d+)?)\s*/\s*(\d+(?:\.\d+)?)\s*", v) if isinstance(v, str) else None
    if m:
        v = float(m[1]) / float(m[2])
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        ops = {">=": v.__ge__, "<=": v.__le__, ">": v.__gt__, "<": v.__lt__, "=": v.__eq__}
        for op, thr, label in SECTION.findall(cell.number_format or ""):
            if ops[op](float(thr)):
                return label
        rest = re.findall(r'(?:^|;)\s*"([^"]+)"\s*$', cell.number_format or "")
        return rest[-1] if rest else None
    return v


def attendance(xl):
    ws = openpyxl.load_workbook(xl)["Attendance"]
    h = [c.value for c in ws[1]]
    days = [x for x in h if str(x)[:2] == "20"]
    rows = {}
    for r in ws.iter_rows(min_row=2):
        if r[1].value:
            rows[str(r[1].value).lower()] = (r[0].value, [mark(r[h.index(d)]) for d in days])
    return days, rows


pct = lambda m: (2 * m.count("Present") + m.count("Late")) / (2 * len(m)) * 100
cert = lambda A, P: "اجتياز" if A > 65 and 0.7 * P + 0.3 * A > 65 else ("حضور" if A > 65 else "بدون")


def table(a):
    best = projects(a.course, a.grades)
    _, att = attendance(a.xlsx)
    out = ["| # | chat name | project score | remaining | notes |", "|---|---|---|---|---|"]
    npass = 0
    for i, (P, d, x, g) in enumerate(sorted(best.values(), key=lambda t: t[2].get("posted", "")), 1):
        e = str(x["student"]).lower()
        A = pct(att[e][1]) if e in att else 0
        if A <= 65:
            rem, note = "—", f"attendance {A:.0f}%"
        elif cert(A, P) == "اجتياز":
            rem, note = "0", ""
            npass += 1
        else:
            rem, note = str(math.floor((65 - 0.3 * A) / 0.7) + 1 - P), g.get("missing", "")
        name = x.get("chat_name") or (att.get(e, ("",))[0])
        out.append(f"| {i} | {name} | {P} | {rem} | {note} |")
    open(a.out, "w", encoding="utf-8").write("\n".join(out) + "\n")
    print(f"submitted {len(best)} | pass {npass} -> {a.out}")


def whatif(a):
    D = a.course
    cp = load(os.path.join(D, "classpulse_marks.json"), {})
    cfg = load(os.path.join(D, "course.json"), {})
    ratio = {p["name"]: {d: (v["answered"], v["total"]) for d, v in p["days"].items()} for p in cp.get("people", [])}
    ov = {e.lower(): v for e, v in cfg.get("status_overrides", {}).items()}
    days, att = attendance(a.xlsx)
    best = projects(D, a.grades)
    res = []
    for e, (name, old) in att.items():
        new = list(old)
        for i, d in enumerate(days):
            q = ratio.get(name, {}).get(d)
            if str(i + 1) in ov.get(e, {}) or not q or old[i] not in ("Present", "Late", "Absent"):
                continue
            r = q[0] / q[1]
            new[i] = "Present" if r >= a.present_min else "Late" if r >= a.late_min else "Absent"
        if new != old:
            P = best.get(e, (0,))[0]
            res.append((name, best.get(e, (None,))[0], pct(old), pct(new), cert(pct(old), P), cert(pct(new), P),
                        [f"{d[5:]} {o}->{n}" for d, o, n in zip(days, old, new) if o != n]))
    res.sort(key=lambda r: (r[4] == r[5], -(r[1] or -1)))
    with open(a.out, "w", encoding="utf-8") as f:
        for r in res:
            f.write(f"{r[0]} | project {r[1]} | attendance {r[2]:.0f}% -> {r[3]:.0f}% | {r[4]} -> {r[5]} | {', '.join(r[6])}\n")
    ch = [r for r in res if r[4] != r[5]]
    print(f"present_min {a.present_min}: marks change for {len(res)} | certificate changes {len(ch)} "
          f"(new pass {sum(1 for r in ch if r[5] == 'اجتياز')}) -> {a.out}")


ap = argparse.ArgumentParser()
ap.add_argument("cmd", choices=["table", "whatif"])
ap.add_argument("course")
ap.add_argument("xlsx")
ap.add_argument("grades")
ap.add_argument("out")
ap.add_argument("--present-min", type=float, default=0.40)
ap.add_argument("--late-min", type=float, default=0.25)
a = ap.parse_args()
{"table": table, "whatif": whatif}[a.cmd](a)
