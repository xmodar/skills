"""Keep a day's Webex chat complete: archive each raw export and rebuild chat-<date>.xlsx as the union of all of them.

    python -I chat_day.py <course_dir> <YYYY-MM-DD> [--since HH:MM] [--out <file>]

Why: Webex re-exports (a) may use the new 12-column format with time-only stamps ("10:03"), which build_workbook.py
cannot parse, and (b) DROP THE OLDEST MESSAGES once the chat grows — a learner whose only message was early loses
attendance. Run this after every new export lands, before --check / build:
- a raw export (more than 3 columns or time-only stamps) is copied to <course_dir>/backup/chat-<date>.raw-HHMMSS.xlsx;
- chat-<date>.xlsx is rewritten as timestamp|user|text, the union of every raw export of that day. For each
  (timestamp, text) the rows come from the NEWEST export containing it, so senders the instructor renamed in a later
  export (two people with the same display name) take the corrected name;
- with --since, messages from that time on (all senders) go to --out (default <course_dir>/chat_new.txt), UTF-8.
Prints one summary line.
"""
import argparse, glob, os, shutil, time, warnings
import openpyxl

warnings.filterwarnings("ignore")


def rows(path, day):
    ws = openpyxl.load_workbook(path).active
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r and r[0] and r[1]:                    # rows without a sender (system lines) are dropped
            t = str(r[0])
            out.append(((f"{day} {t}" if len(t) <= 5 else t[:16]), r[1], r[2]))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("course")
    ap.add_argument("day")
    ap.add_argument("--since")
    ap.add_argument("--out")
    a = ap.parse_args()
    f = os.path.join(a.course, f"chat-{a.day}.xlsx")
    bdir = os.path.join(a.course, "backup")
    os.makedirs(bdir, exist_ok=True)
    archived = ""
    if os.path.exists(f):
        ws = openpyxl.load_workbook(f).active
        first = next(ws.iter_rows(min_row=2, max_row=2, values_only=True), (None,))
        if ws.max_column > 3 or (first[0] and len(str(first[0])) <= 5):
            dst = os.path.join(bdir, f"chat-{a.day}.raw-{time.strftime('%H%M%S')}.xlsx")
            shutil.copy2(f, dst)
            archived = os.path.basename(dst)
    raws = sorted(glob.glob(os.path.join(bdir, f"chat-{a.day}.raw-*.xlsx")), key=os.path.getmtime)
    if not raws:
        print(f"chat-{a.day}.xlsx: no raw exports archived (already in timestamp|user|text form?)")
        return
    groups = {}
    for p in raws:                                  # oldest → newest: a newer export replaces the whole group
        g = {}
        for t, u, x in rows(p, a.day):
            g.setdefault((t, str(x).strip()), []).append((t, u, x))
        groups.update(g)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = a.day
    ws.append(["timestamp", "user", "text"])
    allrows = [r for k in sorted(groups, key=lambda k: k[0]) for r in groups[k]]
    for r in allrows:
        ws.append(list(r))
    wb.save(f)
    newest = len(rows(raws[-1], a.day))
    msg = f"chat-{a.day}.xlsx: {len(allrows)} messages from {len(raws)} raw exports (newest has {newest})"
    if archived:
        msg += f" | archived {archived}"
    if a.since:
        out = a.out or os.path.join(a.course, "chat_new.txt")
        cut = f"{a.day} {a.since}"
        new = [r for r in allrows if r[0] >= cut]
        with open(out, "w", encoding="utf-8") as o:
            for t, u, x in new:
                o.write(f"{t[11:]} | {u} | {x}\n")
        msg += f" | {len(new)} since {a.since} -> {out}"
    print(msg)


if __name__ == "__main__":
    main()
