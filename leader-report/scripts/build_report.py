# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0", "pillow>=10"]
# ///
"""Build the weekly leaders report (اجتماع قادة قطاع بناء القدرات) from a JSON file.

Usage:
  uv run build_report.py data.json --check              # validate + layout check + summary, no file written
  uv run build_report.py data.json -o report.pptx       # generate the deck
  uv run build_report.py --list-icons                   # icon names usable in "icon"

The template (assets/template.pptx) is never modified. The output keeps the title
slide, one slide per department (plus an optional challenges slide after it) and the
closing slide. The template's hidden helper slides are removed, and the script
refuses to write a deck that still contains a hidden slide.
"""
from __future__ import annotations

import argparse
import copy
import io
import json
import math
import re
import sys
from datetime import date
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation

SKILL_DIR = Path(__file__).resolve().parent.parent
TEMPLATE = SKILL_DIR / "assets" / "template.pptx"

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
A = "{%s}" % NS["a"]
R_EMBED = "{%s}embed" % NS["r"]

EMU = 914400  # per inch

# Template slide indexes (0-based)
S_TITLE, S_GUIDE, S_DEPT, S_THANKS, S_EXAMPLE, S_ELEMENTS = range(6)

# Layout geometry (EMU), taken from the template's department slide
TABLE_X = 731520           # rows start at 0.80"
TABLE_W = 9765792          # task rows width (bar column excluded)
FULL_W = 10332720          # header width, rows + bar column
ROWS_TOP = 2048256         # 2.24", right under the header
MAX_BOTTOM = int(6.62 * EMU)
BAR_X = 10533888
BAR_W = 493776
GAP = 109728               # 0.12" gap before the not-started section / between card rows
CARD_H = 512064
CARD_RIGHT_X = 5705856
CARD_LEFT_X = 731520
ICON_CX = 10259568         # icon centre in a task row
CARD_ICON_CX = 10177272    # icon centre in the right card
ROW_ICON_MAX = 292608
CARD_ICON_MAX = 347472
LINE_H = 0.19              # inches per wrapped text line (15pt, 93% spacing)
ROW_H1 = 0.42              # one-line row

NAVY, GREEN = "202B5D", "067E6E"
GROUP_COLORS = {"navy": NAVY, "green": GREEN, "كحلي": NAVY, "أخضر": GREEN}
READY_GREEN = "1E9E62"
READY_WORD = "جاهزة"

MONTHS = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو",
          "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]
WEEKDAYS = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]
STATUS = {"done": "done", "مكتمل": "done", "in_progress": "in_progress", "جار العمل": "in_progress",
          "not_started": "not_started", "لم يبدأ": "not_started"}
ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")


class ReportError(Exception):
    pass


# ---------------------------------------------------------------- formatting

def fmt_date(value, meeting_year: int) -> str:
    """'2026-10-22' -> '22 أكتوبر' (year added when it differs from the meeting year)."""
    if not isinstance(value, str):
        raise ReportError(f"date must be text, got {value!r}")
    value = value.strip()
    if ISO.match(value):
        d = date.fromisoformat(value)
        s = f"{d.day} {MONTHS[d.month - 1]}"
        return s if d.year == meeting_year else f"{s} {d.year}"
    return value


def fmt_meeting_date(value: str) -> str:
    if ISO.match(value.strip()):
        d = date.fromisoformat(value.strip())
        return f"{WEEKDAYS[d.weekday()]} {d.day} {MONTHS[d.month - 1]} {d.year}"
    return value


def est_lines(text: str, width_in: float, pt: float) -> int:
    """Rough line count for Markazi Text (≈0.073" per character at 15pt)."""
    char_w = 0.073 * pt / 15
    per_line = max(1, int(width_in / char_w))
    return sum(max(1, math.ceil(len(part) / per_line)) for part in str(text).split("\n"))


def row_height(lines: int) -> int:
    return int((ROW_H1 + LINE_H * (max(1, lines) - 1)) * EMU)


# ---------------------------------------------------------------- xml helpers

def q(el, path):
    return etree._Element.xpath(el, path, namespaces=NS)


def named(el, name):
    """Direct-or-nested child shape by cNvPr name."""
    res = q(el, f".//*[local-name()='cNvPr' and @name='{name}']/../..")
    return res


def find_shape(slide, name):
    res = named(slide.shapes._spTree, name)
    if not res:
        raise ReportError(f"template shape not found: {name}")
    return res[0]


def xfrm(el):
    """Top-level xfrm of a shape, picture, connector or group."""
    for path in ("./p:spPr/a:xfrm", "./p:grpSpPr/a:xfrm"):
        r = q(el, path)
        if r:
            return r[0]
    raise ReportError("shape without xfrm")


def get_geom(el):
    x = xfrm(el)
    off, ext = x.find(A + "off"), x.find(A + "ext")
    return int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))


def set_geom(el, x=None, y=None, cx=None, cy=None):
    """Move/resize a non-group shape."""
    xf = xfrm(el)
    off, ext = xf.find(A + "off"), xf.find(A + "ext")
    if x is not None:
        off.set("x", str(int(x)))
    if y is not None:
        off.set("y", str(int(y)))
    if cx is not None:
        ext.set("cx", str(int(cx)))
    if cy is not None:
        ext.set("cy", str(int(cy)))


def shift(el, dx=0, dy=0):
    """Move a shape or a whole group subtree (off and chOff alike)."""
    for tag in ("off", "chOff"):
        for o in el.iter(A + tag):
            o.set("x", str(int(o.get("x")) + int(dx)))
            o.set("y", str(int(o.get("y")) + int(dy)))


def set_group_box(grp, y, cy):
    xf = xfrm(grp)
    for tag in ("off", "chOff"):
        xf.find(A + tag).set("y", str(int(y)))
    for tag in ("ext", "chExt"):
        xf.find(A + tag).set("cy", str(int(cy)))


def set_fill(el, rgb):
    for c in q(el, "./p:spPr/a:solidFill/a:srgbClr"):
        c.set("val", rgb)


def set_text(sp, lines, color=None, size=None):
    """Replace the text of a shape, one paragraph per line, keeping the first run's style."""
    if isinstance(lines, str):
        lines = lines.split("\n")
    txBody = q(sp, "./p:txBody")[0]
    paras = txBody.findall(A + "p")
    proto = paras[0]
    for p in paras:
        txBody.remove(p)
    for line in lines:
        p = copy.deepcopy(proto)
        runs = p.findall(A + "r")
        for extra in runs[1:]:
            p.remove(extra)
        for f in p.findall(A + "endParaRPr"):
            p.remove(f)
        r = runs[0]
        r.find(A + "t").text = line
        rpr = r.find(A + "rPr")
        if size:
            rpr.set("sz", str(size))
        if color:
            for c in rpr.iter(A + "srgbClr"):
                c.set("val", color)
        txBody.append(p)


def renumber_ids(slide):
    n = 2
    for c in slide.shapes._spTree.iter():
        if etree.QName(c).localname == "cNvPr":
            c.set("id", str(n))
            n += 1


# ---------------------------------------------------------------- template parts

class Parts:
    """Prototype shapes harvested from the template slides."""

    def __init__(self, prs):
        s = prs.slides
        dept, ex, el = s[S_DEPT], s[S_EXAMPLE], s[S_ELEMENTS]
        self.title = copy.deepcopy(find_shape(dept, "Title 1"))
        self.header = copy.deepcopy(find_shape(dept, "رأس الجدول"))
        self.row = copy.deepcopy(find_shape(dept, "صف مهمة 2"))           # 60% arc row
        self.row_top = get_geom(self.row)[1]
        self.ring_full = copy.deepcopy(named(find_shape(dept, "صف مهمة 1"), "الحلقة مكتملة")[0])
        self.desc_para = copy.deepcopy(q(named(find_shape(dept, "صف مهمة 3"), "المهمة")[0], ".//a:p")[1])
        self.bar = copy.deepcopy(find_shape(dept, "شريط المجموعة اسم المجموعة"))
        self.bar_ns = copy.deepcopy(find_shape(dept, "شريط المجموعة لم تبدأ"))
        self.card = copy.deepcopy(find_shape(dept, "بطاقة لم تبدأ 1"))
        details = copy.deepcopy(find_shape(ex, "تفاصيل قطاعات سماي ٢"))
        self.chip = {
            "done": copy.deepcopy(named(details, "قطاع التعليم")[0]),
            "in_progress": copy.deepcopy(named(details, "قطاع العدل")[0]),
            "not_started": copy.deepcopy(named(details, "قطاع الرياضة")[0]),
        }
        for chip in q(details, "./p:sp[starts-with(p:nvSpPr/p:cNvPr/@name, 'قطاع')]"):
            details.remove(chip)
        self.details = details
        self.details_top = get_geom(details)[1]
        self.icons = self._icons(el, [dept, ex])

    @staticmethod
    def _icons(elements_slide, others):
        """name -> (png bytes, cx, cy). Arabic label and English file stem both work."""
        icons = {}
        items = list(elements_slide.shapes)
        for i, sh in enumerate(items):
            if sh.shape_type != 13:
                continue
            stem = sh._element.nvPicPr.cNvPr.get("descr", "").removesuffix(".png")
            entry = (sh.image.blob, sh.width, sh.height)
            icons.setdefault(stem, entry)
            nxt = items[i + 1] if i + 1 < len(items) else None
            if nxt is not None and nxt.name.startswith("اسم الأيقونة "):
                icons.setdefault(nxt.name.removeprefix("اسم الأيقونة "), entry)
        for slide in others:  # icons that only exist in their faint form (e.g. review_bug)
            stack = list(slide.shapes)
            while stack:
                sh = stack.pop()
                if sh.shape_type == 6:
                    stack.extend(sh.shapes)
                elif sh.shape_type == 13:
                    stem = sh._element.nvPicPr.cNvPr.get("descr", "").removesuffix(".png")
                    icons.setdefault(stem.removesuffix("__faint"), (sh.image.blob, sh.width, sh.height))
        return icons

    def icon(self, name):
        key = (name or "مهام").strip()
        if key not in self.icons:
            raise ReportError(f"unknown icon '{key}'. Run --list-icons for the list.")
        return self.icons[key]


def faint_png(blob):
    """Grey version used on not-started cards (same as the template's __faint icons)."""
    im = Image.open(io.BytesIO(blob)).convert("RGBA")
    grey = Image.new("RGBA", im.size, (166, 166, 166, 255))
    grey.putalpha(im.getchannel("A"))
    out = io.BytesIO()
    grey.save(out, "PNG")
    return out.getvalue()


def add_picture(slide, proto_pic, blob, cx_center, cy_center, max_dim, w, h):
    """Clone a template <p:pic>, point it at `blob`, size it to fit max_dim keeping aspect."""
    pic = copy.deepcopy(proto_pic)
    _, rid = slide.part.get_or_add_image_part(io.BytesIO(blob))
    q(pic, ".//a:blip")[0].set(R_EMBED, rid)
    scale = max_dim / max(w, h)
    cx, cy = w * scale, h * scale
    set_geom(pic, x=cx_center - cx / 2, y=cy_center - cy / 2, cx=cx, cy=cy)
    return pic


# ---------------------------------------------------------------- validation

def require(obj, key, where):
    v = obj.get(key)
    if v is None or (isinstance(v, str) and not v.strip()) or (isinstance(v, list) and not v):
        raise ReportError(f"{where}: '{key}' is missing")
    return v


def normalise(data):
    """Validate the JSON and return (model, warnings). Raises ReportError on blocking problems."""
    warnings = []
    md = require(data, "meeting_date", "report")
    year = date.fromisoformat(md).year if ISO.match(md) else date.today().year
    depts = require(data, "departments", "report")
    model = {"meeting_date": fmt_meeting_date(md), "departments": []}
    for di, d in enumerate(depts, 1):
        dname = require(d, "name", f"department {di}")
        where = f"«{dname}»"
        groups = d.get("groups") or []
        nstart = d.get("not_started") or []
        if not groups and not nstart:
            raise ReportError(f"{where}: no tasks and no not-started items")
        g_out, n_tasks = [], 0
        for gi, g in enumerate(groups):
            tasks = require(g, "tasks", f"{where} group {gi + 1}")
            t_out = []
            for ti, t in enumerate(tasks, 1):
                tw = f"{where} task {ti} of group «{g.get('name', '')}»"
                name = require(t, "name", tw)
                prog = t.get("progress")
                if not isinstance(prog, (int, float)) or not 0 <= prog <= 100:
                    raise ReportError(f"{tw} «{name}»: progress must be a number 0-100")
                prog = int(round(prog))
                done = require(t, "done", f"{tw} «{name}»")
                nxt = require(t, "next", f"{tw} «{name}»")
                ready = t.get("readiness")
                if prog == 100:
                    if ready and ready != READY_WORD:
                        warnings.append(f"{tw} «{name}» is 100%: readiness set to «{READY_WORD}» (was {ready!r})")
                    ready_lines = [READY_WORD]
                else:
                    if not ready:
                        raise ReportError(f"{tw} «{name}»: 'readiness' date is missing")
                    if ready == READY_WORD:
                        warnings.append(f"{tw} «{name}» is «{READY_WORD}» but only {prog}%")
                    ready_lines = [fmt_date(x, year) for x in (ready if isinstance(ready, list) else [ready])]
                    if any(re.search(r"(أسبوع|أيام|يوم|شهر|أسابيع)", x) for x in ready_lines):
                        warnings.append(f"{tw} «{name}»: readiness looks like a duration, use a date")
                details = []
                for x in t.get("details") or []:
                    st = STATUS.get(x.get("status"))
                    if not st:
                        raise ReportError(f"{tw} «{name}»: detail {x!r} needs status done/in_progress/not_started")
                    details.append({"label": require(x, "label", tw), "status": st})
                t_out.append({"name": name, "description": t.get("description", ""), "icon": t.get("icon"),
                              "progress": prog, "done": done, "next": nxt, "ready": ready_lines,
                              "details": details})
            n_tasks += len(t_out)
            g_out.append({"name": g.get("name", ""), "color": g.get("color"), "tasks": t_out})
        if n_tasks > 6:
            warnings.append(f"{where}: {n_tasks} tasks in the table (the template rule is at most 6)")
        n_out = []
        for ni, n in enumerate(nstart, 1):
            nw = f"{where} not-started item {ni}"
            n_out.append({"name": require(n, "name", nw), "icon": n.get("icon"),
                          "start": fmt_date(require(n, "start", nw), year),
                          "ready": fmt_date(require(n, "ready", nw), year)})
        c_out = []
        for ci, c in enumerate(d.get("challenges") or [], 1):
            cw = f"{where} challenge {ci}"
            c_out.append({"challenge": require(c, "challenge", cw), "impact": require(c, "impact", cw),
                          "support": require(c, "support", cw), "due": fmt_date(require(c, "due", cw), year)})
        model["departments"].append({"name": dname, "groups": g_out, "not_started": n_out, "challenges": c_out})
    return model, warnings


# ---------------------------------------------------------------- layout

def task_lines(t):
    name_lines = est_lines(t["name"], 2.40, 16) + (est_lines(t["description"], 2.40, 15) if t["description"] else 0)
    return max(name_lines,
               est_lines(t["done"], 2.10, 15),
               est_lines(t["next"], 2.25, 15),
               sum(est_lines(x, 1.45, 16) for x in t["ready"]))


def chip_width(label):
    return 163622 + 50661 * len(label)


CHIP_RIGHT, CHIP_GAP, CHIP_MIN_X, CHIP_PITCH = 7699248, 90485, 804672, 292608


def chip_layout(details):
    """[(x, line, width)] laid out right-to-left, wrapping onto extra lines."""
    out, x_right, line = [], CHIP_RIGHT, 0
    for d in details:
        w = chip_width(d["label"])
        if x_right - w < CHIP_MIN_X and x_right != CHIP_RIGHT:
            line, x_right = line + 1, CHIP_RIGHT
        out.append((x_right - w, line, w))
        x_right -= w + CHIP_GAP
    return out


def details_height(details):
    lines = 1 + max(l for _, l, _ in chip_layout(details))
    return 365760 + (lines - 1) * CHIP_PITCH


def plan_department(dept):
    """Vertical plan of a department slide; returns (items, bottom)."""
    y, items = ROWS_TOP, []
    for gi, g in enumerate(dept["groups"]):
        g_top = y
        for t in g["tasks"]:
            h = row_height(task_lines(t))
            items.append(("row", t, y, h))
            y += h
            if t["details"]:
                dh = details_height(t["details"])
                items.append(("details", t, y, dh))
                y += dh
        if g["name"]:
            items.append(("bar", g, g_top, y - g_top, gi))
    if dept["not_started"]:
        if dept["groups"]:
            y += GAP
        n_rows = math.ceil(len(dept["not_started"]) / 2)
        top = y
        for i, n in enumerate(dept["not_started"]):
            r = i // 2
            items.append(("card", n, top + r * (CARD_H + GAP), i % 2))
        y = top + n_rows * CARD_H + (n_rows - 1) * GAP
        items.append(("bar_ns", None, top, y - top))
    return items, y


def challenge_lines(c):
    return max(est_lines(c["challenge"], 3.00, 16), est_lines(c["impact"], 2.55, 15),
               est_lines(c["support"], 3.10, 15), est_lines(c["due"], 1.45, 16))


def plan_challenges(dept):
    y, items = ROWS_TOP, []
    for c in dept["challenges"]:
        h = row_height(challenge_lines(c))
        items.append((c, y, h))
        y += h
    return items, y


def check_fit(model):
    problems = []
    for d in model["departments"]:
        _, bottom = plan_department(d)
        if bottom > MAX_BOTTOM:
            over = (bottom - MAX_BOTTOM) / EMU
            problems.append(f"«{d['name']}» does not fit on one slide (over by {over:.2f}\"). "
                            "Shorten the longest texts, drop or merge tasks, or remove details rows.")
        if d["challenges"]:
            _, cb = plan_challenges(d)
            if cb > MAX_BOTTOM:
                problems.append(f"«{d['name']}» challenges do not fit on one slide; shorten or split them.")
    return problems


# ---------------------------------------------------------------- builders

def new_slide_like(prs, src):
    slide = prs.slides.add_slide(src.slide_layout)
    tree = slide.shapes._spTree
    for sh in list(slide.shapes):
        tree.remove(sh._element)
    return slide


def build_row(slide, parts, t, y, h):
    row = copy.deepcopy(parts.row)
    set_group_box(row, y, h)
    centre_dy = (y + h / 2) - (parts.row_top + get_geom(parts.row)[3] / 2)
    for el in list(row):
        if etree.QName(el).localname not in ("sp", "cxnSp", "pic", "grpSp"):
            continue
        name = q(el, "./*[1]/p:cNvPr/@name")[0]
        if name in ("خلفية الصف", "المهمة", "ما تم", "الخطوة القادمة"):
            set_geom(el, y=y, cy=h)
        elif name == "فاصل عمود":
            set_geom(el, y=y + 54864, cy=h - 109728)
        elif name == "الجاهزية":
            set_geom(el, y=y + 45720, cy=h - 91440)
        elif etree.QName(el).localname == "pic":
            blob, w, hh = parts.icon(t["icon"])
            new = add_picture(slide, el, blob, ICON_CX, y + h / 2, ROW_ICON_MAX, w, hh)
            row.replace(el, new)
        else:  # ring group, check mark, arrow: keep vertically centred
            shift(el, dy=centre_dy)

    # task name (+ optional description paragraph)
    task = named(row, "المهمة")[0]
    set_text(task, [t["name"]])
    if t["description"]:
        p = copy.deepcopy(parts.desc_para)
        q(p, ".//a:t")[0].text = t["description"]
        q(task, "./p:txBody")[0].append(p)
    set_text(named(row, "ما تم")[0], t["done"])
    set_text(named(row, "الخطوة القادمة")[0], t["next"])

    ready = named(row, "الجاهزية")[0]
    is_ready = t["ready"] == [READY_WORD]
    set_text(ready, t["ready"], color=READY_GREEN if is_ready else NAVY,
             size=1400 if len(t["ready"]) > 1 else None)

    # progress ring
    ring = q(row, "./p:grpSp[starts-with(p:nvGrpSpPr/p:cNvPr/@name, 'حلقة الإنجاز')]")[0]
    q(ring, "./p:nvGrpSpPr/p:cNvPr")[0].set("name", f"حلقة الإنجاز {t['progress']}%")
    arc = named(ring, "قوس الإنجاز (اسحب المقبض الأصفر)")[0]
    pct = named(ring, "النسبة")[0]
    prog = t["progress"]
    if prog >= 100:
        full = copy.deepcopy(parts.ring_full)
        ax, ay, _, _ = get_geom(arc)
        fx, fy, _, _ = get_geom(full)
        shift(full, ax - fx, ay - fy)
        ring.replace(arc, full)
        set_text(pct, "100%", color=READY_GREEN)
    elif prog <= 0:
        ring.remove(arc)
        set_text(pct, "0%", color="595959")
    else:
        end = (270 + prog * 3.6) % 360
        for gd in q(arc, ".//a:gd[@name='adj2']"):
            gd.set("fmla", f"val {int(round(end * 60000))}")
        set_text(pct, f"{prog}%")
    slide.shapes._spTree.append(row)


def build_details(slide, parts, t, y, h):
    grp = copy.deepcopy(parts.details)
    shift(grp, dy=y - parts.details_top)
    set_group_box(grp, y, h)
    set_geom(named(grp, "خلفية التفاصيل")[0], cy=h)
    for d, (x, line, w) in zip(t["details"], chip_layout(t["details"])):
        chip = copy.deepcopy(parts.chip[d["status"]])
        set_geom(chip, x=x, y=y + 64008 + line * CHIP_PITCH, cx=w)
        q(chip, "./p:nvSpPr/p:cNvPr")[0].set("name", f"قطاع {d['label']}")
        set_text(chip, d["label"])
        grp.append(chip)
    slide.shapes._spTree.append(grp)


def corner_adj(w, h, radius=54864):
    return int(radius * 100000 / max(1, min(w, h)))


def build_bar(slide, proto, label, y, h, color=None):
    bar = copy.deepcopy(proto)
    by, bh = y + 27432, h - 54864
    set_geom(bar, x=BAR_X, y=by, cx=BAR_W, cy=bh)
    for gd in q(bar, ".//a:prstGeom/a:avLst/a:gd[@name='adj']"):
        gd.set("fmla", f"val {corner_adj(BAR_W, bh)}")
    if color:
        set_fill(bar, color)
    words = [label] if len(label) <= 8 else label.split()
    longest = max(len(w) for w in words)
    size = max(800, min(13, 117 // longest) * 100)   # the 0.54" bar holds ~9 chars at 13pt
    if len(words) * size / 100 * 0.9 / 72 * EMU > bh:   # lines do not fit vertically
        size = 1100
        if len(words) * 11 * 0.9 / 72 * EMU > bh:
            words = [label]
    set_text(bar, words, size=size)
    q(bar, "./p:nvSpPr/p:cNvPr")[0].set("name", f"شريط المجموعة {label}")
    slide.shapes._spTree.append(bar)


def build_card(slide, parts, n, y, col):
    card = copy.deepcopy(parts.card)
    cx0, cy0, _, _ = get_geom(card)
    dx = 0 if col == 0 else CARD_LEFT_X - CARD_RIGHT_X
    shift(card, dx=dx, dy=y - cy0)
    pic = q(card, "./p:pic")[0]
    blob, w, h = parts.icon(n["icon"])
    new = add_picture(slide, pic, faint_png(blob), CARD_ICON_CX + dx, y + CARD_H / 2, CARD_ICON_MAX, w, h)
    card.replace(pic, new)
    txt = named(card, "نص البطاقة")[0]
    paras = q(txt, "./p:txBody/a:p")
    q(paras[0], ".//a:t")[0].text = n["name"]
    runs = q(paras[1], "./a:r")
    for extra in runs[1:]:
        paras[1].remove(extra)
    q(runs[0], "./a:t")[0].text = f"يبدأ {n['start']}  ·  الجاهزية {n['ready']}"
    slide.shapes._spTree.append(card)


def build_department(prs, parts, dept):
    slide = new_slide_like(prs, prs.slides[S_DEPT])
    tree = slide.shapes._spTree
    title = copy.deepcopy(parts.title)
    set_text(title, dept["name"])
    tree.append(title)
    tree.append(copy.deepcopy(parts.header))
    items, _ = plan_department(dept)
    gi_color = 0
    for it in items:
        kind = it[0]
        if kind == "row":
            build_row(slide, parts, it[1], it[2], it[3])
        elif kind == "details":
            build_details(slide, parts, it[1], it[2], it[3])
        elif kind == "bar":
            g, gi = it[1], it[4]
            color = GROUP_COLORS.get((g.get("color") or "").strip()) or (NAVY if gi_color % 2 == 0 else GREEN)
            gi_color += 1
            build_bar(slide, parts.bar, g["name"], it[2], it[3], color)
        elif kind == "card":
            build_card(slide, parts, it[1], it[2], it[3])
        elif kind == "bar_ns":
            build_bar(slide, parts.bar_ns, "لم تبدأ", it[2], it[3])
    renumber_ids(slide)
    return slide


# Challenges slide columns, right to left: (header text, x, width)
CH_COLS = [("التحدي", 8.40, 3.55), ("الأثر", 5.75, 2.50), ("الدعم المطلوب من القادة", 2.30, 3.30), ("الموعد", 0.86, 1.30)]
CH_SEPS = [8.30, 5.65, 2.22]


def build_challenges(prs, parts, dept):
    slide = new_slide_like(prs, prs.slides[S_DEPT])
    tree = slide.shapes._spTree
    title = copy.deepcopy(parts.title)
    set_text(title, f"{dept['name']}: التحديات والدعم المطلوب")
    tree.append(title)

    header = copy.deepcopy(parts.header)
    cols = [named(header, n)[0] for n in ("عمود المهمة", "عمود الإنجاز", "عمود ما تم", "عمود الخطوة القادمة", "عمود الجاهزية")]
    header.remove(cols.pop(1))   # no progress column
    for box, (label, x, w) in zip(cols, CH_COLS):
        set_geom(box, x=x * EMU, cx=w * EMU)
        set_text(box, label)
    tree.append(header)

    risk_blob, rw, rh = parts.icon("مخاطر")
    items, _ = plan_challenges(dept)
    for c, y, h in items:
        row = copy.deepcopy(parts.row)
        set_group_box(row, y, h)
        xf = xfrm(row)
        for tag in ("ext", "chExt"):
            xf.find(A + tag).set("cx", str(FULL_W))
        for el in list(row):
            kind = etree.QName(el).localname
            if kind not in ("sp", "cxnSp", "pic", "grpSp"):
                continue
            name = q(el, "./*[1]/p:cNvPr/@name")[0]
            if kind == "grpSp" or name in ("علامة تم", "سهم"):
                row.remove(el)
            elif name == "خلفية الصف":
                set_geom(el, y=y, cy=h, cx=FULL_W)
            elif name == "فاصل عمود":
                row.remove(el)
            elif kind == "pic":
                row.replace(el, add_picture(slide, el, risk_blob, 11.72 * EMU, y + h / 2, ROW_ICON_MAX, rw, rh))
        seps_proto = q(parts.row, "./p:cxnSp")[0]
        for sx in CH_SEPS:
            sep = copy.deepcopy(seps_proto)
            set_geom(sep, x=sx * EMU, y=y + 54864, cy=h - 109728)
            row.append(sep)
        boxes = {"المهمة": ("challenge", 0, 3.00), "ما تم": ("impact", 1, None),
                 "الخطوة القادمة": ("support", 2, None)}
        for name, (key, col, width) in boxes.items():
            box = named(row, name)[0]
            _, x, w = CH_COLS[col]
            set_geom(box, x=x * EMU, y=y, cx=(width or w) * EMU, cy=h)
            set_text(box, c[key])
        due = named(row, "الجاهزية")[0]
        set_geom(due, y=y + 45720, cy=h - 91440)
        set_text(due, c["due"], color=NAVY)
        tree.append(row)
    renumber_ids(slide)
    return slide


def build(model, out_path):
    prs = Presentation(str(TEMPLATE))
    parts = Parts(prs)
    # title slide date
    title_slide = prs.slides[S_TITLE]
    for sh in title_slide.shapes:
        if sh.has_text_frame and "[تاريخ الاجتماع]" in sh.text_frame.text:
            set_text(sh._element, model["meeting_date"])

    order = [prs.slides[S_TITLE]]
    for d in model["departments"]:
        order.append(build_department(prs, parts, d))
        if d["challenges"]:
            order.append(build_challenges(prs, parts, d))
    order.append(prs.slides[S_THANKS])

    order_ids = [s.slide_id for s in order]
    sld_lst = prs.slides._sldIdLst
    by_id = {int(e.get("id")): e for e in sld_lst}
    for e in list(sld_lst):
        if int(e.get("id")) not in order_ids:
            prs.part.drop_rel(e.get(f"{{{NS['r']}}}id"))
        sld_lst.remove(e)
    for sid in order_ids:
        sld_lst.append(by_id[sid])

    for s in prs.slides:
        s._element.attrib.pop("show", None)
    prs.save(out_path)
    verify(out_path, len(order))


def verify(path, expected):
    prs = Presentation(str(path))
    hidden = [i for i, s in enumerate(prs.slides, 1) if s._element.get("show") in ("0", "false")]
    if hidden:
        raise ReportError(f"hidden slides found in output: {hidden}")
    if len(prs.slides) != expected:
        raise ReportError(f"expected {expected} slides, found {len(prs.slides)}")
    leftovers = []
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes._spTree.iter(A + "t"):
            if sh.text and re.search(r"\[[^\]]*\]", sh.text):
                leftovers.append(f"slide {i}: {sh.text}")
    if leftovers:
        raise ReportError("template placeholders left in output: " + "; ".join(leftovers))


# ---------------------------------------------------------------- summary

def summary(model):
    out = [f"تاريخ الاجتماع: {model['meeting_date']}", ""]
    slide_no = 2
    for d in model["departments"]:
        out.append(f"شريحة {slide_no}: {d['name']}")
        slide_no += 1
        for g in d["groups"]:
            out.append(f"  المجموعة: {g['name'] or '(بدون شريط)'}")
            for t in g["tasks"]:
                desc = f" ({t['description']})" if t["description"] else ""
                out.append(f"    - {t['name']}{desc} | {t['progress']}% | ما تم: {t['done']} | "
                           f"القادمة: {t['next']} | الجاهزية: {' / '.join(t['ready'])} | أيقونة: {t['icon'] or 'مهام'}")
                if t["details"]:
                    lab = {"done": "مكتمل", "in_progress": "جار العمل", "not_started": "لم يبدأ"}
                    out.append("      تفاصيل: " + "، ".join(f"{x['label']} ({lab[x['status']]})" for x in t["details"]))
        if d["not_started"]:
            out.append("  لم تبدأ:")
            for n in d["not_started"]:
                out.append(f"    - {n['name']} | يبدأ {n['start']} | الجاهزية {n['ready']} | أيقونة: {n['icon'] or 'مهام'}")
        if d["challenges"]:
            out.append(f"شريحة {slide_no}: {d['name']}: التحديات والدعم المطلوب")
            slide_no += 1
            for c in d["challenges"]:
                out.append(f"    - {c['challenge']} | الأثر: {c['impact']} | الدعم: {c['support']} | الموعد: {c['due']}")
        out.append("")
    out.append(f"شريحة {slide_no}: شكراً لكم")
    out.append(f"إجمالي الشرائح: {slide_no} (بدون شرائح مخفية)")
    return "\n".join(out)


def main():
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("data", nargs="?", help="report JSON")
    ap.add_argument("-o", "--output", help="output .pptx path")
    ap.add_argument("--check", action="store_true", help="validate and print the summary only")
    ap.add_argument("--list-icons", action="store_true")
    args = ap.parse_args()

    if args.list_icons:
        parts = Parts(Presentation(str(TEMPLATE)))
        arabic = sorted(k for k in parts.icons if not re.match(r"^[a-z_]+$", k))
        english = sorted(k for k in parts.icons if re.match(r"^[a-z_]+$", k))
        print("Arabic names:", "، ".join(arabic))
        print("File names:", ", ".join(english))
        return 0
    if not args.data:
        ap.error("data JSON is required")
    try:
        data = json.loads(Path(args.data).read_text(encoding="utf-8"))
        model, warnings = normalise(data)
        parts = Parts(Presentation(str(TEMPLATE)))
        for d in model["departments"]:   # fail early on unknown icons
            for g in d["groups"]:
                for t in g["tasks"]:
                    parts.icon(t["icon"])
            for n in d["not_started"]:
                parts.icon(n["icon"])
        problems = check_fit(model)
        for w in warnings:
            print("WARNING:", w)
        if problems:
            for p in problems:
                print("ERROR:", p)
            return 2
        if args.check:
            print(summary(model))
            print("\nOK: data is complete and fits the template.")
            return 0
        if not args.output:
            ap.error("--output is required unless --check is used")
        out = Path(args.output)
        out.parent.mkdir(parents=True, exist_ok=True)
        build(model, out)
        print(f"OK: wrote {out} ({sum(1 + bool(d['challenges']) for d in model['departments']) + 2} slides, none hidden)")
        return 0
    except ReportError as e:
        print("ERROR:", e)
        return 2


if __name__ == "__main__":
    sys.exit(main())
