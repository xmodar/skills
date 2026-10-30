"""Build the single-sheet audit workbook.

Usage: python build_workbook.py results.json output.xlsx

results.json:
{
  "title": "...", "max_score": 20, "caveat": "...",
  "parts": [{"key": "Q1a", "max": 5}, ...],
  "policy": {"threshold": 0.3},
  "students": [{
    "id": "4414...", "submission": "attempts/files",
    "scores": {"Q1a": 3.5, ...},
    "raw_notes": "...",
    "ai_signs": "...", "ai_pct": 20, "ai_conf": 60,
    "cross_signs": "...", "cross_partners": "...", "cross_pct": 10, "cross_conf": 60,
    "reasoning": "..."
  }]
}
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864")
GROUPS = {"Raw evaluation": "2F5597", "AI signs": "7F6000", "Cross-student": "833C0B", "Outcome": "375623"}
INPUT = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def main(src, dst):
    data = json.load(open(src, encoding="utf-8"))
    parts = data["parts"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit"

    ws["A1"] = data["title"]
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = data["caveat"]
    ws["A2"].font = Font(name=FONT, italic=True, size=10, color="C00000")
    ws["A3"] = "Action threshold (evidence weight = % x confidence; below this, no deduction)"
    ws["A3"].font = Font(name=FONT, size=10)
    ws["A4"] = "Maximum score"
    ws["A4"].font = Font(name=FONT, size=10)
    n_fixed = 2 + len(parts)
    ws.cell(3, n_fixed + 1, data["policy"]["threshold"]).number_format = "0%"
    ws.cell(4, n_fixed + 1, data["max_score"])
    for row in (3, 4):
        cell = ws.cell(row, n_fixed + 1)
        cell.fill, cell.border = INPUT, BOX
        cell.font = Font(name=FONT, size=10, color="0000FF", bold=True)
    thresh = f"${L(n_fixed + 1)}$3"

    cols = [("Student ID", 13, None), ("Submission", 20, None)]
    cols += [(f"{p['key']} (/{p['max']})", 8, "Raw evaluation") for p in parts]
    cols += [("Raw total", 9, "Raw evaluation"), ("Grading notes", 46, "Raw evaluation"),
             ("AI: signs observed", 46, "AI signs"), ("AI cheating %", 9, "AI signs"), ("AI confidence %", 10, "AI signs"),
             ("Cross: evidence", 46, "Cross-student"), ("Cross: suspected partners", 16, "Cross-student"),
             ("Cross cheating %", 9, "Cross-student"), ("Cross confidence %", 10, "Cross-student"),
             ("Reasoning and notes", 70, "Outcome"), ("Penalty %", 9, "Outcome"), ("Cheating-deducted score", 11, "Outcome")]
    idx = {name: i + 1 for i, (name, _, _) in enumerate(cols)}
    group_row, head_row = 6, 7
    for i, (name, width, group) in enumerate(cols, 1):
        ws.column_dimensions[L(i)].width = width
        cell = ws.cell(head_row, i, name)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=GROUPS.get(group, "1F3864")) if group else HEAD
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.border = BOX
    ws.row_dimensions[head_row].height = 44
    start = None
    for i, (_, _, group) in enumerate(cols + [("", 0, "__end")], 1):
        if start is None or group != cols[start - 1][2]:
            if start and cols[start - 1][2]:
                ws.merge_cells(start_row=group_row, start_column=start, end_row=group_row, end_column=i - 1)
                c = ws.cell(group_row, start, cols[start - 1][2])
                c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
                c.fill = PatternFill("solid", fgColor=GROUPS[cols[start - 1][2]])
                c.alignment = Alignment(horizontal="center")
            start = i

    first = head_row + 1
    for r, s in enumerate(data["students"], first):
        ws.cell(r, 1, s["id"])
        ws.cell(r, 2, s["submission"])
        for j, p in enumerate(parts, 3):
            ws.cell(r, j, s["scores"].get(p["key"], 0))
        raw_col = idx["Raw total"]
        ws.cell(r, raw_col, f"=SUM({L(3)}{r}:{L(2 + len(parts))}{r})")
        ws.cell(r, idx["Grading notes"], s["raw_notes"])
        ws.cell(r, idx["AI: signs observed"], s["ai_signs"])
        ws.cell(r, idx["AI cheating %"], s["ai_pct"])
        ws.cell(r, idx["AI confidence %"], s["ai_conf"])
        ws.cell(r, idx["Cross: evidence"], s["cross_signs"])
        ws.cell(r, idx["Cross: suspected partners"], s["cross_partners"])
        ws.cell(r, idx["Cross cheating %"], s["cross_pct"])
        ws.cell(r, idx["Cross confidence %"], s["cross_conf"])
        ws.cell(r, idx["Reasoning and notes"], s["reasoning"])
        a, ac = L(idx["AI cheating %"]), L(idx["AI confidence %"])
        x, xc = L(idx["Cross cheating %"]), L(idx["Cross confidence %"])
        weight = f"MAX({a}{r}*{ac}{r},{x}{r}*{xc}{r})/10000"
        pen = ws.cell(r, idx["Penalty %"], f"=IF({weight}>={thresh},{weight},0)")
        pen.number_format = "0%"
        final = ws.cell(r, idx["Cheating-deducted score"], f"=ROUND({L(raw_col)}{r}*(1-{L(idx['Penalty %'])}{r}),2)")
        final.font = Font(name=FONT, bold=True, size=10)
        for i in range(1, len(cols) + 1):
            c = ws.cell(r, i)
            c.border = BOX
            if c is not final:
                c.font = Font(name=FONT, size=10)
            wrap = i in {2, idx["Grading notes"], idx["AI: signs observed"], idx["Cross: evidence"],
                         idx["Cross: suspected partners"], idx["Reasoning and notes"]}
            c.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=None if wrap else "center")
        longest = max(len(str(s[k])) / w for k, w in (("raw_notes", 46), ("ai_signs", 46), ("cross_signs", 46), ("reasoning", 70)))
        ws.row_dimensions[r].height = min(409, max(30, 13.5 * (longest * 1.15 + 1)))

    last = first + len(data["students"]) - 1
    red, amber = PatternFill("solid", bgColor="F8CBAD"), PatternFill("solid", bgColor="FFE699")
    for name in ("AI cheating %", "Cross cheating %"):
        rng = f"{L(idx[name])}{first}:{L(idx[name])}{last}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["61"], fill=red))
        ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["36", "60"], fill=amber))
    ws.freeze_panes = ws.cell(first, 2)
    ws.auto_filter.ref = f"A{head_row}:{L(len(cols))}{last}"
    wb.calculation.fullCalcOnLoad = True
    wb.save(dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
